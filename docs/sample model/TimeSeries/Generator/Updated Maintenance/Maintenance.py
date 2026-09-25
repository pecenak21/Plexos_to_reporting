# -*- coding: utf-8 -*-
"""
Created on Tue Feb  10 10:16:48 2026

@author: Z20045
"""

import re
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

# ---------------- CONFIG ----------------
INPUT_FILE = "Maintenance_input.xlsx"
OUTPUT_FILE = "Maintenance_updated.csv"

START_DATE = "2026-01-01"
END_DATE = "2044-12-31"

UNIT_COL = "Unit"
START_COL = "Start date"
END_COL = "End date"
# ----------------------------------------

RANGE_RE = re.compile(r"^(?P<prefix>.*?)(?P<start>\d+)\s*-\s*(?P<end>\d+)\s*$")


def expand_unit(unit: str) -> list[str]:
   
    unit = str(unit).strip()
    m = RANGE_RE.match(unit)
    if not m:
        return [unit]

    prefix = m.group("prefix")
    start_s = m.group("start")
    end_s = m.group("end")

    start = int(start_s)
    end = int(end_s)
    step = 1 if end >= start else -1
    width = len(start_s)

    return [f"{prefix}{str(n).zfill(width)}".strip()
            for n in range(start, end + step, step)]


def format_output_excel(path: str):
    wb = load_workbook(path)
    ws = wb.active

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    ws.freeze_panes = "A2"

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        header = ws.cell(row=1, column=col_idx).value
        ws.column_dimensions[col_letter].width = (
            10 if header in ("Year", "Month", "Day") else max(14, len(str(header)) + 2)
        )

    wb.save(path)


def main():
    df = pd.read_excel(INPUT_FILE)

    # Validate expected columns
    for col in (UNIT_COL, START_COL, END_COL):
        if col not in df.columns:
            raise ValueError(
                f"Input must have columns: {UNIT_COL}, {START_COL}, {END_COL}. "
                f"Missing '{col}'. Found: {list(df.columns)}"
            )

    # Parse dates
    df[START_COL] = pd.to_datetime(df[START_COL], errors="coerce").dt.normalize()
    df[END_COL] = pd.to_datetime(df[END_COL], errors="coerce").dt.normalize()
    df = df.dropna(subset=[UNIT_COL, START_COL, END_COL])

    # Ensure start <= end
    swap = df[START_COL] > df[END_COL]
    if swap.any():
        df.loc[swap, [START_COL, END_COL]] = df.loc[swap, [END_COL, START_COL]].values

    
    rows = []
    for _, r in df.iterrows():
        for u in expand_unit(r[UNIT_COL]):
            rows.append((u, r[START_COL], r[END_COL]))

    df_exp = pd.DataFrame(rows, columns=["unit", "start", "end"])

  
    days = pd.date_range(START_DATE, END_DATE, freq="D")
    days_np = days.values.astype("datetime64[D]")  # daily granularity for clean comparisons

    out = pd.DataFrame({
        "Year": days.year,
        "Month": days.month,
        "Day": days.day
    })

   
    units = sorted(df_exp["unit"].unique().tolist())
    for u in units:
        out[u] = 0

    
    df_exp["start_np"] = df_exp["start"].values.astype("datetime64[D]")
    df_exp["end_np"] = df_exp["end"].values.astype("datetime64[D]")


    for u in units:
        windows = df_exp[df_exp["unit"] == u][["start_np", "end_np"]].to_numpy()
        mask = np.zeros(len(days_np), dtype=bool)

        for s_np, e_np in windows:
            mask |= (days_np >= s_np) & (days_np <= e_np)

        out[u] = mask.astype(int)

  
   
    out.to_csv("Maintenance_updated.csv", index=False)
   

    print(f"Wrote: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
